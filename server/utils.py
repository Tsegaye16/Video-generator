import os
import aiohttp

import io
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from config import settings
import google.generativeai as genai
from models import SlideData,  TableData
from typing import Any,  List
import logging


logger = logging.getLogger(__name__)

def get_slide_count(file_path: str) -> int:
    try:
        prs = Presentation(file_path)
        return len(prs.slides)
    except PackageNotFoundError:
        raise ValueError("Cannot process legacy .ppt format reliably for slide count. Please use .pptx.")
    except Exception as e:
        raise ValueError(f"Error processing PowerPoint file: {e}")




def summarize_table(table_data: TableData) -> str:
    if not table_data or not table_data.rows:
        return ""
    header = ", ".join(table_data.rows[0])
    num_rows = len(table_data.rows) - 1
    return f"Table with header '{header}' and {num_rows} data rows."

def format_slide_content_for_llm(slide: SlideData, extracted_content_path: str) -> List[Any]:
    parts = []
    if slide.title:
        parts.append(f"Slide Title: {slide.title}\n")
    if slide.text_elements:
        parts.append("Slide Text:\n" + "\n".join([f"- {text}" for text in slide.text_elements]) + "\n")
    if slide.tables:
        table_summaries = [summarize_table(table) for table in slide.tables]
        parts.append("Slide Tables Summary:\n" + "\n".join([f"- {summary}" for summary in table_summaries]) + "\n")
    for img_info in slide.images:
        img_path = os.path.join(extracted_content_path, img_info.filename)
        if os.path.exists(img_path):
            try:
                img = Image.open(img_path)
                mime_type = None
                if img.format == "PNG":
                    mime_type = "image/png"
                elif img.format == "JPEG":
                    mime_type = "image/jpeg"
                elif img.format == "GIF":
                    mime_type = "image/gif"
                elif img.format == "BMP":
                    mime_type = "image/bmp"
                elif img_info.content_type:
                    mime_type = img_info.content_type
                if mime_type:
                    img_bytes = io.BytesIO()
                    img.save(img_bytes, format=img.format)
                    parts.append(genai.Part.from_data(img_bytes.getvalue(), mime_type=mime_type))
                    parts.append(f"\nImage Description (from file: {img_info.filename}): ")
                else:
                    pass
                   
            except Exception as e:
                pass
        else:
            pass
    return parts

async def merge_with_logo(background_image_data: bytes, logo_url: str, output_format: str = "PNG") -> bytes:
    try:
        # Open the background image
        background = Image.open(BytesIO(background_image_data)).convert("RGBA")
        logger.debug(f"Background image opened: {background.size}, mode: {background.mode}")

        # Download the logo from the provided URL
        async with aiohttp.ClientSession() as session:
            async with session.get(logo_url) as response:
                if response.status != 200:
                    logger.error(f"Failed to download logo from {logo_url}: {response.status}")
                    raise Exception(f"Failed to download logo: {response.status}")
                logo_data = await response.read()
                logger.debug(f"Logo downloaded: {len(logo_data)} bytes")

        # Open and process the logo
        logo = Image.open(BytesIO(logo_data)).convert("RGBA")
        logo = logo.resize((90, 90), Image.LANCZOS)  # Use LANCZOS for high-quality resizing
        logger.debug(f"Logo resized: {logo.size}, mode: {logo.mode}")

        # Merge the images
        x = background.width - logo.width  # Top-right corner
        y = 0
        composite = Image.new("RGBA", background.size)
        composite.paste(background, (0, 0))
        composite.paste(logo, (x, y), logo)  # Use logo as mask for transparency
        logger.debug(f"Images merged: composite size {composite.size}")

        # Save the merged image to a buffer
        merged_buffer = BytesIO()
        composite.save(merged_buffer, format=output_format)
        merged_buffer.seek(0)
        result = merged_buffer.getvalue()
        logger.info(f"Merged image created: {len(result)} bytes, format: {output_format}")
        return result

    except Exception as e:
        logger.error(f"Error merging with logo: {e}")
        return None


